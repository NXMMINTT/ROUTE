"""
อ่านไฟล์ที่ผู้ใช้อัปโหลด (.vrp / .txt / .csv / .xlsx) แล้วคืนเนื้อหารูปแบบ CVRPLIB ให้ parse_vrp อ่านต่อ
ตาราง (.csv / .xlsx) แปลงเป็นข้อความ CVRPLIB ก่อน solver และคลัง Benchmark จึงไม่ต้องรู้จักรูปแบบอื่น

รูปแบบตาราง: แถวแรกเป็นหัวคอลัมน์ แถวข้อมูลแถวแรก = คลังสินค้า (depot)
- x, y, demand : จำเป็น
- capacity     : จำเป็น ใส่ค่าในแถวใดก็ได้ (ปกติใส่แถวแรกแถวเดียว)
- id           : ไม่ใส่ = เลขตามลำดับแถว 1, 2, 3, ...
- name, vehicles : ไม่บังคับ (ชื่อ instance / จำนวนรถ)
"""

import csv
import io
import math
import os

from app.services.solver import VrpFormatError

TABLE_EXTENSIONS = {".csv", ".xlsx"}
REQUIRED = ("x", "y", "demand")


def read_upload(filename: str, raw: bytes) -> str:
    """ไฟล์ที่อัปโหลด → ข้อความ CVRPLIB (นามสกุลอื่นนอกจาก .csv / .xlsx ถือเป็นข้อความ CVRPLIB อยู่แล้ว)"""
    ext = os.path.splitext(filename or "")[1].lower()
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    if ext == ".csv":
        return table_to_vrp(_csv_rows(raw), stem)
    if ext == ".xlsx":
        return table_to_vrp(_xlsx_rows(raw), stem)
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        raise VrpFormatError("ไฟล์ .vrp / .txt ต้องเป็นไฟล์ข้อความ (UTF-8)")


def _csv_rows(raw: bytes) -> list[list]:
    try:
        text = raw.decode("utf-8-sig")  # utf-8-sig: ตัด BOM ที่ Excel ใส่ไว้หน้าไฟล์ CSV
    except UnicodeDecodeError:
        raise VrpFormatError("ไฟล์ CSV ต้องบันทึกเป็น UTF-8 (ใน Excel เลือก CSV UTF-8)")
    return list(csv.reader(io.StringIO(text)))


def _xlsx_rows(raw: bytes) -> list[list]:
    from openpyxl import load_workbook  # import เฉพาะตอนใช้: ไม่ให้ทั้ง backend พังถ้ายังไม่ได้ติดตั้ง

    try:
        sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
    except Exception:
        raise VrpFormatError("อ่านไฟล์ Excel ไม่ได้ ตรวจสอบว่าเป็นไฟล์ .xlsx")
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _is_blank(v) -> bool:
    return v is None or str(v).strip() == ""


def _number(v, column: str, row: int) -> float:
    try:
        n = float(str(v).strip())
    except ValueError:
        raise VrpFormatError(f"แถว {row}: คอลัมน์ {column} ต้องเป็นตัวเลข (พบ \"{v}\")")
    if not math.isfinite(n):
        raise VrpFormatError(f"แถว {row}: คอลัมน์ {column} ต้องเป็นตัวเลข (พบ \"{v}\")")
    return n


def _integer(v, column: str, row: int) -> int:
    n = _number(v, column, row)
    if n != int(n):
        raise VrpFormatError(f"แถว {row}: คอลัมน์ {column} ต้องเป็นจำนวนเต็ม (พบ \"{v}\")")
    return int(n)


def table_to_vrp(rows: list[list], fallback_name: str) -> str:
    rows = [r for r in rows if not all(_is_blank(v) for v in r)]  # ข้ามแถวว่าง
    if len(rows) < 3:
        raise VrpFormatError("ตารางต้องมีหัวคอลัมน์ + คลังสินค้า 1 แถว + ลูกค้าอย่างน้อย 1 แถว")

    header = [str(h).strip().lower() if not _is_blank(h) else "" for h in rows[0]]
    col = {h: i for i, h in enumerate(header) if h}
    missing = [c for c in (*REQUIRED, "capacity") if c not in col]
    if missing:
        raise VrpFormatError(f"ไม่พบคอลัมน์ {', '.join(missing)} (ต้องมี x, y, demand, capacity)")

    def cell(r, name):
        i = col.get(name)
        return r[i] if i is not None and i < len(r) else None

    def first(name):
        """ค่าแรกที่ไม่ว่างของคอลัมน์ (capacity / name / vehicles ใส่แถวเดียวพอ)"""
        return next((cell(r, name) for r in rows[1:] if not _is_blank(cell(r, name))), None)

    data = rows[1:]
    nodes = []
    for n, r in enumerate(data, start=1):
        line = n + 1  # เลขแถวตามที่ผู้ใช้เห็นในไฟล์ (แถว 1 = หัวคอลัมน์)
        for c in REQUIRED:
            if _is_blank(cell(r, c)):
                raise VrpFormatError(f"แถว {line}: คอลัมน์ {c} ว่าง")
        nid = _integer(cell(r, "id"), "id", line) if not _is_blank(cell(r, "id")) else n
        nodes.append((nid, _number(cell(r, "x"), "x", line), _number(cell(r, "y"), "y", line),
                      _integer(cell(r, "demand"), "demand", line)))

    ids = [n[0] for n in nodes]
    if len(set(ids)) != len(ids) or min(ids) <= 0:
        raise VrpFormatError("คอลัมน์ id ต้องเป็นจำนวนเต็มบวกและไม่ซ้ำกัน")
    capacity = first("capacity")
    if capacity is None:
        raise VrpFormatError("คอลัมน์ capacity ว่าง ใส่ความจุรถอย่างน้อยในแถวแรก")
    capacity = _integer(capacity, "capacity", 2)

    name = first("name")
    name = str(name).strip() if name is not None else fallback_name
    vehicles = first("vehicles")
    comment = f"No of trucks: {_integer(vehicles, 'vehicles', 2)}" if vehicles is not None else ""

    lines = [
        f"NAME : {name}",
        f"COMMENT : ({comment})",
        "TYPE : CVRP",
        f"DIMENSION : {len(nodes)}",
        "EDGE_WEIGHT_TYPE : EXACT_2D",
        f"CAPACITY : {capacity}",
        "NODE_COORD_SECTION",
        *(f"{nid} {x!r} {y!r}" for nid, x, y, _ in nodes),
        "DEMAND_SECTION",
        *(f"{nid} {d}" for nid, _, _, d in nodes),
        "DEPOT_SECTION",
        f"{nodes[0][0]}",
        "-1",
        "EOF",
    ]
    return "\n".join(lines) + "\n"
